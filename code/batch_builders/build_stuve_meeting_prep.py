from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("outputs")
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "Dr_Stuve_Meeting_Prep_2026-08-20.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203040"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GOLD = "FFF4D6"
GOLD = "8A6400"
LIGHT_GRAY = "F2F4F7"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D0D5DD", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_field(paragraph, code):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = code
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_numbering_definition(doc, num_id, abstract_id, fmt, text, left=270, hanging=270):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left + hanging))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left + hanging))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)


def add_bullet(doc, text, bold_prefix=None, num_id=101):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text.split(":", 1)
        set_run_font(p.add_run(first + ":"), bold=True)
        set_run_font(p.add_run(rest))
    else:
        set_run_font(p.add_run(text))
    return p


def add_numbered(doc, text, num_id=102):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_body(doc, text, bold_prefix=None, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text.split(":", 1)
        set_run_font(p.add_run(first + ":"), bold=True)
        set_run_font(p.add_run(rest), italic=italic)
    else:
        set_run_font(p.add_run(text), italic=italic)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(label + "  "), size=10.5, color=label_color, bold=True)
    set_run_font(p.add_run(text), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, [1700, 7660])
    set_table_borders(table, color="D0D5DD", size="4")
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        set_cell_width(cells[0], 1700)
        set_cell_width(cells[1], 7660)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_shading(cells[0], LIGHT_BLUE)
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        set_run_font(p0.add_run(label), size=10, color=DARK_BLUE, bold=True)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.1
        set_run_font(p1.add_run(value), size=10, color=INK)
    set_table_geometry(table, [1700, 7660])
    if table.rows:
        set_repeat_table_header(table.rows[0])
    return table


def add_evidence_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    set_repeat_table_header(table.rows[0])
    headers = ("Evidence", "Current result", "Correct interpretation")
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], DARK_BLUE)
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=9.5, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    for label, result, meaning in rows:
        cells = table.add_row().cells
        for idx, text in enumerate((label, result, meaning)):
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[idx], "F8FAFC")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(text), size=9.1, color=INK, bold=(idx == 0))
    set_table_geometry(table, [2100, 3180, 4080])
    set_table_borders(table, color="CBD5E1", size="4")
    return table


def setup_styles(doc):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, before, after, color in (
        (1, 16, 18, 10, BLUE),
        (2, 13, 14, 7, BLUE),
        (3, 12, 10, 5, DARK_BLUE),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    add_numbering_definition(doc, 101, 101, "bullet", "•", 269, 270)
    add_numbering_definition(doc, 102, 102, "decimal", "%1.", 269, 270)
    add_numbering_definition(doc, 103, 103, "bullet", "☐", 269, 270)



def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("MEETING CALL CARD"), size=9.5, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run("Dr. Olaf Stüve + Dr. Marco Tapia Maltos"), size=24, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("EBV–MS HLA-DR15 computational pMHC project  |  Thursday, August 20, 2026"), size=11.5, color=MUTED, italic=True)

    table = doc.add_table(rows=1, cols=3)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [3120, 3120, 3120])
    set_table_borders(table, color="D8E2EE", size="4")
    metrics = [
        ("LOGISTICS", "Later-afternoon time + link pending"),
        ("TARGET OUTCOME", "One decisive next validation step"),
        ("MEETING STYLE", "20 minutes • concise • decision-focused"),
    ]
    for i, (label, value) in enumerate(metrics):
        cell = table.cell(0, i)
        set_cell_shading(cell, "F6F9FC")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(label), size=8, color=BLUE, bold=True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=9.2, color=INK, bold=True)


def build_document():
    doc = Document()
    setup_styles(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "CONFIRMATION CHECK",
        "The latest Gmail message says Dr. Tapia Maltos is checking whether a later-afternoon time works for Dr. Stüve. The proposed 3:45 p.m. time conflicts with school, and no calendar event or meeting link is present yet. Verify the final time and link before the call.",
        fill=LIGHT_GOLD,
        label_color=GOLD,
    )

    add_heading(doc, "The one outcome you want", 1)
    add_body(
        doc,
        "Get their expert judgment on whether the primary EBNA1–MBP HLA-DRB1*15:01 pMHC lead is biologically worth experimental follow-up—and identify the single most decisive next validation step.",
    )

    add_heading(doc, "Your 60-second opening", 1)
    add_callout(
        doc,
        "SAY THIS",
        "Thank you both for meeting with me. I built a reproducible computational workflow to ask whether EBV and CNS peptides can adopt similar surfaces when presented in equivalent positions by HLA-DR15. I froze the candidate pairs and register rules, aligned the HLA grooves, compared the same P1–P9 positions, and then tested the strongest leads against score-blind matched controls. One EBNA1–MBP pair remained structurally consistent: its median exposed-position difference was 0.643 Å versus 7.964 Å for three strict controls. I know this does not prove natural presentation, shared-TCR recognition, activation, or an MS mechanism. I would like your judgment on whether the hypothesis is biologically coherent and what experiment would test it most directly.",
    )

    add_heading(doc, "If they ask for the project in one sentence", 2)
    add_body(
        doc,
        "A register-aware computational screen nominated one HLA-DRB1*15:01 EBNA1–MBP pMHC pair as a testable structural-resemblance hypothesis, while explicitly separating pMHC geometry from TCR cross-reactivity and disease claims.",
        italic=True,
    )

    doc.add_page_break()

    add_heading(doc, "Evidence snapshot", 1)
    add_body(
        doc,
        "Use their expertise deliberately: Dr. Stüve for clinical/neuroimmunology relevance and claim boundaries; Dr. Tapia Maltos for practical biological critique and realistic assay design. Because Dr. Stüve is near a grant deadline, lead with the result, caveat, and decision question.",
        after=5,
    )
    add_key_value_table(
        doc,
        [
            ("Primary lead", "EBNA1 residues 482–496 (AEGLRALLARSHVER) vs MBP residues 245–263 (LSRFSWGAEGQRPGFGYGG)."),
            ("Predicted cores", "LRALLARSH vs WGAEGQRPG; both inputs recorded in an HLA-DRB1*15:01 context. These are computational register hypotheses, not experimentally mapped registers."),
            ("Primary result", "Target median exposed-position RMSD 0.643 Å; equal-weight strict-background median 7.964 Å; background-minus-target difference 7.321 Å."),
            ("Robustness", "All four target job-pair medians were below the background median. Leave-one-control-out differences were 5.261–10.086 Å. Technical bootstrap interval: 1.620–12.759 Å."),
            ("Control limit", "Three frozen strict controls; exploratory rank 1 of 4. The 0.25 empirical tail fraction is the minimum possible here and is not a p-value."),
        ],
    )

    add_heading(doc, "What each evidence layer means", 2)
    add_evidence_table(
        doc,
        [
            ("Rank 1 lead", "Consistent against its three strict primary controls.", "The pair deserves experimental prioritization—not a biological conclusion."),
            ("Rank 2", "Triplex capsid protein 1–PLP1; technical interval crosses zero.", "Sensitivity-only; do not present as a second primary lead."),
            ("BALF5–MBP", "Established cross-reactive calibration system across DRB5*01:01 and DRB1*15:01.", "Useful precedent and calibration, but not same-allele validation of the new lead."),
            ("TCR modeling", "Five-chain Ob.1A12 model failed to recover the known experimental TCR orientation.", "Excluded from evidence; no binding or cross-recognition claim."),
        ],
    )

    add_heading(doc, "Non-negotiable claim boundary", 2)
    add_callout(
        doc,
        "SAFE CLAIM",
        "The analysis prioritizes a testable pMHC structural-resemblance hypothesis. It does not establish natural peptide presentation, shared-TCR binding, T-cell activation, cross-reactivity, molecular mimicry in patients, or an EBV-driven MS mechanism.",
        fill=LIGHT_BLUE,
    )

    doc.add_page_break()

    add_heading(doc, "Questions worth spending the meeting on", 1)
    questions = [
        "Does the EBNA1 482–496 / MBP 245–263 pairing look biologically plausible after considering antigen processing, natural presentation, and relevant immune compartments—or is there an obvious biological reason to deprioritize it?",
        "Is HLA-DRB1*15:01 the right primary framing for this new pair, and how should DRB5*01:01 and the broader DR15 haplotype be handled without conflating the alleles?",
        "What is the fastest falsification-first experiment, and what minimal controls or comparison groups would make it convincing?",
        "Is ‘register-aware computational pMHC candidate prioritization’ the right paper/science-fair frame, or is there a more clinically meaningful but still honest way to state the contribution?",
        "What single flaw would make you reject the hypothesis now? If it is worth pursuing, could you recommend someone with HLA class-II register or clone-defined T-cell assay expertise to review a one-page protocol?",
    ]
    for q in questions:
        add_numbered(doc, q)

    add_heading(doc, "Suggested 20-minute flow", 2)
    add_key_value_table(
        doc,
        [
            ("0–2 min", "Thank them; give the 60-second opening; state the one outcome you want."),
            ("2–6 min", "Show only the primary lead and explain the three-control limitation."),
            ("6–16 min", "Ask biological-plausibility and falsification-first questions; take exact notes."),
            ("16–20 min", "Repeat back the decision, confirm the next deliverable, and ask permission for one focused follow-up."),
        ],
    )

    add_heading(doc, "Close with a concrete ask", 2)
    add_callout(
        doc,
        "CLOSING",
        "Based on what you have seen, what is the single most informative next experiment or analysis? If I turn your recommendation into a one-page protocol, would you or someone you recommend be willing to review it?",
        fill=LIGHT_GOLD,
        label_color=GOLD,
    )

    add_heading(doc, "Notes to capture verbatim", 2)
    for text in (
        "Most important biological criticism:",
        "Recommended first experiment and success/failure criterion:",
        "Contact, next deliverable, and follow-up date:",
    ):
        p = doc.add_paragraph()
        apply_numbering(p, 103)
        set_run_font(p.add_run(text + "  ______________________________________________"), size=10.2)

    doc.core_properties.title = "Dr. Stüve Meeting Prep — EBV–MS HLA-DR15 Project"
    doc.core_properties.subject = "Decision-focused call card for August 20, 2026"
    doc.core_properties.author = "Anish Sharma"
    doc.core_properties.keywords = "EBV, multiple sclerosis, HLA-DR15, pMHC, meeting prep"
    doc.save(OUT_PATH)
    print(OUT_PATH.resolve())


if __name__ == "__main__":
    build_document()
