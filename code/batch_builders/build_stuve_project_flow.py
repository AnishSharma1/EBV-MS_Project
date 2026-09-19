from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("outputs")
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "EBV_MS_Project_Flow_for_Dr_Stuve_2026-08-20.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "203040"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
LIGHT_GOLD = "FFF4D6"
GOLD = "8A6400"
GREEN = "237A57"
WHITE = "FFFFFF"


def set_run_font(run, size=11, color=INK, bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_numbering_definition(doc, num_id, abstract_id, fmt="decimal", text="%1."):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), fmt)
    level.append(number_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), text)
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_props = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_props.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_props.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_props.append(spacing)
    level.append(p_props)
    abstract.append(level)
    numbering.append(abstract)

    instance = OxmlElement("w:num")
    instance.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    instance.append(abstract_ref)
    numbering.append(instance)


def apply_numbering(paragraph, num_id):
    p_props = paragraph._p.get_or_add_pPr()
    numbering_props = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), str(num_id))
    numbering_props.append(level)
    numbering_props.append(number_id)
    p_props.append(numbering_props)


def set_repeat_table_header(row):
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for key, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:tcW"))
    if node is None:
        node = OxmlElement("w:tcW")
        props.append(node)
    node.set(qn("w:w"), str(width))
    node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr
    width_node = props.find(qn("w:tblW"))
    if width_node is None:
        width_node = OxmlElement("w:tblW")
        props.append(width_node)
    width_node.set(qn("w:w"), str(sum(widths)))
    width_node.set(qn("w:type"), "dxa")
    indent_node = props.find(qn("w:tblInd"))
    if indent_node is None:
        indent_node = OxmlElement("w:tblInd")
        props.append(indent_node)
    indent_node.set(qn("w:w"), str(indent))
    indent_node.set(qn("w:type"), "dxa")
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D0D5DD", size="4"):
    props = table._tbl.tblPr
    borders = props.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for level, size, color, before, after in (
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    add_numbering_definition(doc, 201, 201)
    add_numbering_definition(doc, 202, 202)
    add_numbering_definition(doc, 203, 203, fmt="bullet", text="•")


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("PROJECT FLOW PLAN"), size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("EBV-MS Molecular Mimicry Project"), size=24, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    set_run_font(
        p.add_run("Completed DRB1*15:01 workflow and proposed multi-allele expansion"),
        size=13,
        color=MUTED,
        italic=True,
    )

    for label, value in (
        ("To", "Dr. Olaf Stüve"),
        ("From", "Anish Sharma"),
        ("Date", "August 20, 2026"),
        ("Purpose", "Summarize the project from its initial question through the proposed next phase."),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        set_run_font(p.add_run(label + ": "), size=10.5, color=INK, bold=True)
        set_run_font(p.add_run(value), size=10.5, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(8)
    props = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    props.append(borders)


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size="4")
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.10
    set_run_font(p.add_run(label + "  "), size=10.5, color=label_color, bold=True)
    set_run_font(p.add_run(text), size=10.5, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_flow_step(doc, num_id, title, status, text):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run(title + "  "), size=10.4, color=DARK_BLUE, bold=True)
    status_color = GREEN if status == "COMPLETED" else GOLD if status == "PROPOSED" else MUTED
    set_run_font(p.add_run("[" + status + "]  "), size=9.2, color=status_color, bold=True)
    set_run_font(p.add_run(text), size=10.4, color=INK)
    return p


def add_bullet(doc, text, num_id=203, bold_prefix=None):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text.split(":", 1)
        set_run_font(p.add_run(first + ":"), size=10.3, bold=True)
        set_run_font(p.add_run(rest), size=10.3)
    else:
        set_run_font(p.add_run(text), size=10.3)
    return p


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        shade(table.rows[0].cells[idx], LIGHT_GRAY)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=9.5, color=DARK_BLUE, bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if row_index % 2 == 1:
                shade(cells[idx], "FAFBFC")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(value), size=9.2, color=INK, bold=(idx == 0))
    set_table_geometry(table, widths)
    set_table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_sources(doc):
    add_heading(doc, "Selected sources", 2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("1. "), size=8.7, color=MUTED, bold=True)
    add_hyperlink(p, "Moutsianas et al., Nature Genetics (2015)", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4874245/")
    set_run_font(p.add_run(" - high-resolution map of class II HLA risk in multiple sclerosis."), size=8.7, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("2. "), size=8.7, color=MUTED, bold=True)
    add_hyperlink(p, "Patsopoulos et al., PLOS Genetics (2013)", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3836799/")
    set_run_font(p.add_run(" - independent HLA effects and peptide-binding-groove variation."), size=8.7, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("3. "), size=8.7, color=MUTED, bold=True)
    add_hyperlink(p, "Abramson et al., Nature (2024)", "https://doi.org/10.1038/s41586-024-07487-w")
    set_run_font(p.add_run(" - AlphaFold 3 methods and model scope."), size=8.7, color=MUTED)


def build_document():
    doc = Document()
    setup_document(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "WORKING OBJECTIVE",
        "Determine whether selected EBV and central-nervous-system peptides can form similar T-cell-facing surfaces when presented by MS-associated HLA class II molecules, then prioritize the most falsifiable candidates for experimental follow-up.",
    )

    add_heading(doc, "Project flow: initial question to current position", 1)
    flow = [
        ("Define the biological question", "COMPLETED", "Ask whether EBV-derived and CNS-derived peptides could present comparable exposed surfaces in an MS-associated HLA context."),
        ("Build an evidence-traceable peptide universe", "COMPLETED", "Curate exact EBV and human peptide sequences, source proteins, coordinates, assay evidence, and HLA restrictions from IEDB and primary literature."),
        ("Lock the DRB1*15:01 register protocol", "COMPLETED", "Retain a predeclared P1-P9 hierarchy, separate experimentally mapped from predicted registers, and preserve unresolved candidates rather than forcing them into the ranking."),
        ("Run register-aware candidate screening", "COMPLETED", "Compare residues only when they occupy equivalent HLA-groove positions and prioritize candidates without selecting favorable windows after seeing results."),
        ("Model three-chain pMHC complexes", "COMPLETED", "Predict mature HLA-DRA, mature HLA-DRB1*15:01, and the exact peptide as separate chains; fit the HLA groove before comparing exposed peptide positions."),
        ("Evaluate against score-blind controls", "COMPLETED", "Use source-validated background peptides selected before structural review and matched by prespecified length and binding-rank rules."),
        ("Define the current technical lead", "COMPLETED", "EBNA1 residues 482-496 versus MBP residues 245-263 showed a median exposed-position RMSD of 0.643 A, compared with 7.964 A for three strict controls. This prioritizes a testable pair; it is not proof of presentation or cross-reactivity."),
        ("Expand to additional MS-associated alleles", "PROPOSED", "Repeat the full evidence, register, control, and structural workflow independently for a provisional panel of DRB1*13:03, DRB1*03:01, and DRB1*08:01."),
        ("Translate computational leads into experiments", "FUTURE", "For the strongest allele-specific candidates, test binding/register, natural presentation, and clone-defined T-cell recognition in that order."),
    ]
    for title, status, text in flow:
        add_flow_step(doc, 201, title, status, text)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "Proposed multi-allele expansion", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.10
    set_run_font(p.add_run("DRB1*15:01 will remain the completed pilot arm. Each new allele will be analyzed independently, and raw structural scores will not be pooled across alleles. The provisional panel reflects class II signals reported by "), size=10.5)
    add_hyperlink(p, "Moutsianas et al. (2015)", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4874245/")
    set_run_font(p.add_run(" and "), size=10.5)
    add_hyperlink(p, "Patsopoulos et al. (2013)", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3836799/")
    set_run_font(p.add_run("."), size=10.5)

    add_matrix(
        doc,
        ("Provisional allele", "Reason for inclusion", "Design caveat"),
        (
            ("HLA-DRB1*13:03", "Strong additive MS-risk association", "Relatively uncommon; ancestry and evidence coverage must be reported"),
            ("HLA-DRB1*03:01", "Replicated class II risk signal", "Risk is largely recessive; effect should not be summarized by one simple rank"),
            ("HLA-DRB1*08:01", "Independent additive risk signal", "Smaller effect; retain as a separate allele-specific test"),
        ),
        (1900, 3400, 4060),
    )

    add_heading(doc, "Ordered next steps", 2)
    next_steps = [
        ("Freeze the allele panel and HLA sequences", "Confirm the three DRB1 alleles using association strength, ancestry context, and peptide evidence; record source-verified mature HLA-DRA and allele-specific DRB1 sequences."),
        ("Build allele-specific peptide universes", "Retrieve exact EBV and CNS peptides for each allele and calculate binding/register hypotheses without using structural results."),
        ("Predeclare candidates and controls", "Select one primary and secondary EBV-CNS pair per allele, plus score-blind controls matched by prespecified length and binding-rank bins."),
        ("Run the 27-job discovery batch", "Model nine pMHC entities per allele with one fixed seed: two EBV peptides, two CNS peptides, three strict controls, and two calibration binders."),
        ("Analyze and run the 30-job robustness batch", "Apply pMHC/register quality checks, compare P2/P3/P5/P7/P8 after groove fitting, then add two seeds for each primary pair and its three controls."),
        ("Prioritize experiments", "Advance only candidates that remain coherent across register, quality, structural-control, and technical-stability checks."),
    ]
    for title, text in next_steps:
        add_flow_step(doc, 202, title, "NEXT", text)

    add_heading(doc, "Questions for guidance", 2)
    for question in (
        "Is the proposed DRB1*13:03 / DRB1*03:01 / DRB1*08:01 panel the most defensible first expansion, or should a different allele or HLA-DQ/DP arm take priority?",
        "Should the next experimental gate prioritize direct peptide-HLA binding/register validation, natural presentation by immunopeptidomics, or clone-defined T-cell recognition?",
        "Would the allele-by-allele replication framework and score-blind control design provide a sufficiently coherent foundation for a computational manuscript?",
    ):
        add_bullet(doc, question)

    add_callout(
        doc,
        "CLAIM BOUNDARY",
        "The workflow prioritizes allele-specific pMHC structural-resemblance hypotheses. It does not establish natural peptide presentation, shared-TCR binding, T-cell activation, cross-reactivity in patients, molecular mimicry, or an EBV-driven mechanism of multiple sclerosis.",
        fill=LIGHT_GOLD,
        label_color=GOLD,
    )

    doc.core_properties.title = "EBV-MS Project Flow and Multi-Allele Expansion"
    doc.core_properties.subject = "Mentor-facing project flow plan"
    doc.core_properties.author = "Anish Sharma"
    doc.core_properties.keywords = "EBV, multiple sclerosis, HLA, pMHC, AlphaFold 3"
    doc.save(OUT_PATH)
    print(OUT_PATH.resolve())


if __name__ == "__main__":
    build_document()
